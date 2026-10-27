"""작성한 서류 → PDF·DOCX (api.md 3.31).

원본 양식 파일이 없는 앱 양식도 있어서, 칸 이름·작성 내용을 표 형태로 정리한 제출용 사본을 만든다.
민감 칸(주민등록번호·계좌·서명)은 비워 두고 "직접 작성"으로, 공무원 기재란은 "담당자 기재"로 표시한다.
"""

import io
from functools import lru_cache

import pymupdf
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

EMPTY_TEXT = {"handwrite": "(양식에 직접 작성)", "office": "(담당자 기재란)"}
FOOTER = "온로드 문서 도우미로 작성한 사본입니다. 제출 전 공식 양식과 내용을 확인해 주세요."

PAGE_W, PAGE_H = pymupdf.paper_size("a4")
MARGIN = 50
LABEL_W = 130
GAP = 10
FONT_SIZE = 10.5
LINE_H = FONT_SIZE * 1.55
PAD = 6
GRAY = (0.45, 0.45, 0.45)
LINE = (0.75, 0.75, 0.75)
FONT_NAME = "kr"


@lru_cache
def _font() -> pymupdf.Font:
    return pymupdf.Font("cjk")  # PyMuPDF 내장 한글 지원 글꼴 (Droid Sans Fallback)


def _wrap(text: str, width: float, size: float) -> list[str]:
    font = _font()
    lines: list[str] = []
    for para in text.split("\n") or [""]:
        line = ""
        for ch in para:
            if font.text_length(line + ch, fontsize=size) > width and line:
                lines.append(line)
                line = ch.lstrip() if ch == " " else ch
            else:
                line += ch
        lines.append(line)
    return lines


def _display(spec: dict, value: str) -> tuple[str, bool]:
    """(보여 줄 글, 회색 여부)"""
    if spec["value_source"] in EMPTY_TEXT:
        return EMPTY_TEXT[spec["value_source"]], True
    return value.strip() or "-", not value.strip()


def render_pdf(title: str, rows: list[tuple[dict, str]]) -> bytes:
    doc = pymupdf.open()
    state = {"page": None, "y": 0.0}

    def new_page() -> None:
        page = doc.new_page(width=PAGE_W, height=PAGE_H)
        page.insert_font(fontname=FONT_NAME, fontbuffer=_font().buffer)
        state["page"], state["y"] = page, MARGIN

    def text_at(x: float, y: float, s: str, size: float = FONT_SIZE, color=(0, 0, 0)) -> None:
        state["page"].insert_text((x, y + size), s, fontname=FONT_NAME, fontsize=size, color=color)

    new_page()
    title_size = 17
    tw = _font().text_length(title, fontsize=title_size)
    text_at((PAGE_W - tw) / 2, state["y"], title, title_size)
    state["y"] += title_size + 22

    value_x = MARGIN + LABEL_W + GAP
    value_w = PAGE_W - MARGIN - value_x
    section = object()
    for spec, value in rows:
        if spec.get("section") != section:
            section = spec.get("section")
            if section:
                if state["y"] + 30 > PAGE_H - MARGIN:
                    new_page()
                state["y"] += 6
                text_at(MARGIN, state["y"], section, 12)
                state["y"] += 12 + 8
        shown, gray = _display(spec, value)
        label_lines = _wrap(spec["label"], LABEL_W - PAD, FONT_SIZE)
        value_lines = _wrap(shown, value_w - PAD * 2, FONT_SIZE)
        bottom = PAGE_H - MARGIN - 20
        n = max(len(label_lines), len(value_lines))
        if state["y"] + min(n, 3) * LINE_H + PAD * 2 > bottom:
            new_page()
        # 긴 칸은 페이지를 넘겨 이어서 그린다
        while label_lines or value_lines:
            fit = max(1, int((bottom - state["y"] - PAD * 2) // LINE_H))
            part_l, label_lines = label_lines[:fit], label_lines[fit:]
            part_v, value_lines = value_lines[:fit], value_lines[fit:]
            height = max(len(part_l), len(part_v)) * LINE_H + PAD * 2
            page, y = state["page"], state["y"]
            page.draw_rect(pymupdf.Rect(MARGIN, y, PAGE_W - MARGIN, y + height), color=LINE, width=0.6)
            page.draw_line((value_x - GAP / 2, y), (value_x - GAP / 2, y + height), color=LINE, width=0.6)
            for i, s in enumerate(part_l):
                text_at(MARGIN + PAD, y + PAD + i * LINE_H, s, color=(0.15, 0.15, 0.15))
            for i, s in enumerate(part_v):
                text_at(value_x + PAD - GAP / 2, y + PAD + i * LINE_H, s, color=GRAY if gray else (0, 0, 0))
            state["y"] += height
            if label_lines or value_lines:
                new_page()

    for i, page in enumerate(doc):
        page.insert_text(
            (MARGIN, PAGE_H - MARGIN / 2), f"{FOOTER}  ({i + 1}/{doc.page_count})",
            fontname=FONT_NAME, fontsize=8, color=GRAY,
        )
    doc.subset_fonts()
    data = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return data


def _korean_font(run, name: str = "맑은 고딕") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def render_docx(title: str, rows: list[tuple[dict, str]]) -> bytes:
    d = DocxDocument()
    head = d.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = head.add_run(title)
    run.bold = True
    run.font.size = Pt(17)
    _korean_font(run)

    section = object()
    table = None
    for spec, value in rows:
        if spec.get("section") != section:
            section = spec.get("section")
            if section:
                p = d.add_paragraph()
                r = p.add_run(section)
                r.bold = True
                r.font.size = Pt(12)
                _korean_font(r)
                table = None
        if table is None:
            table = d.add_table(rows=0, cols=2)
            table.style = "Table Grid"
        shown, gray = _display(spec, value)
        cells = table.add_row().cells
        for cell, text, is_gray, bold in ((cells[0], spec["label"], False, True), (cells[1], shown, gray, False)):
            cell.paragraphs[0].text = ""
            lines = text.split("\n")
            for i, line in enumerate(lines):
                para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                r = para.add_run(line)
                r.bold = bold
                r.font.size = Pt(10.5)
                _korean_font(r)
                if is_gray:
                    r.font.color.rgb = RGBColor(0x73, 0x73, 0x73)

    foot = d.add_paragraph()
    r = foot.add_run(FOOTER)
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x73, 0x73, 0x73)
    _korean_font(r)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
